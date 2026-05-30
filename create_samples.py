"""One-shot script to generate sample resume files in sample_resumes/."""
import os
from pathlib import Path

import docx
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.enums import TA_LEFT

OUT = Path("sample_resumes")
OUT.mkdir(exist_ok=True)


# ── helpers ──────────────────────────────────────────────────────────────────

def make_pdf(path: Path, lines: list[str]) -> None:
    doc = SimpleDocTemplate(str(path), pagesize=LETTER,
                            leftMargin=inch, rightMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()
    heading = ParagraphStyle("H", parent=styles["Heading1"], fontSize=14, spaceAfter=4)
    normal  = ParagraphStyle("N", parent=styles["Normal"],  fontSize=10, leading=14)

    story = []
    for line in lines:
        if line == "":
            story.append(Spacer(1, 6))
        elif line.isupper() and len(line) < 60:
            story.append(Paragraph(line, heading))
        else:
            story.append(Paragraph(line, normal))
    doc.build(story)


def make_docx(path: Path, lines: list[str]) -> None:
    document = docx.Document()
    for line in lines:
        if line == "":
            document.add_paragraph("")
        elif line.isupper() and len(line) < 60:
            document.add_heading(line, level=2)
        else:
            document.add_paragraph(line)
    document.save(str(path))


def make_txt(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


# ── Resume 1: Alice Chen — ML Engineer  (expected score ~8.5) ─────────────

ALICE = [
    "Alice Chen",
    "Senior ML Engineer | alice.chen@email.com | github.com/alicechen-ml",
    "",
    "SUMMARY",
    "5 years of experience designing and deploying production ML systems. Led a team of 4 engineers "
    "building a real-time recommendation engine serving 10M daily users. Contributor to HuggingFace "
    "Transformers. Published 2 papers on efficient fine-tuning of large language models.",
    "",
    "EXPERIENCE",
    "Senior ML Engineer — TechCorp (2022–Present)",
    "- Designed RAG pipeline using LangChain + FAISS reducing hallucination rate by 34%",
    "- Fine-tuned Llama 2 7B on proprietary dataset using LoRA; deployed via vLLM on AWS SageMaker",
    "- Maintained MLflow experiment tracking and built CI/CD for model promotion",
    "- Reduced inference latency 60% via quantization (GPTQ) and dynamic batching",
    "",
    "ML Engineer — DataStart (2020–2022)",
    "- Built XGBoost fraud detection model (AUC 0.97) processing $2B in transactions annually",
    "- Implemented Kubeflow Pipelines for automated retraining on drift detection signals",
    "- Mentored 2 junior engineers in PyTorch and MLOps best practices",
    "",
    "SKILLS",
    "Python, PyTorch, TensorFlow, scikit-learn, HuggingFace Transformers, LangChain, FAISS, "
    "vLLM, MLflow, Kubeflow, AWS SageMaker, Docker, Kubernetes, SQL",
    "",
    "EDUCATION",
    "M.S. Computer Science (ML specialization) — Stanford University, 2020",
    "B.S. Mathematics — UC Berkeley, 2018",
    "",
    "PUBLICATIONS",
    "\"Parameter-Efficient Fine-Tuning of LLMs for Domain Adaptation\" — NeurIPS 2023 Workshop",
    "\"Scalable Vector Search for Production RAG Systems\" — ICLR 2024 Workshop",
]

make_pdf(OUT / "alice_chen_ml_engineer.pdf", ALICE)
print("Created alice_chen_ml_engineer.pdf")


# ── Resume 2: Bob Martinez — DevOps  (expected score ~4.5) ───────────────

BOB = [
    "Bob Martinez",
    "DevOps / Platform Engineer | bob.martinez@email.com",
    "",
    "SUMMARY",
    "8 years in infrastructure, CI/CD, and platform engineering. Recently exploring AI tooling "
    "integration into developer workflows. Comfortable with Python scripting.",
    "",
    "EXPERIENCE",
    "Senior DevOps Engineer — CloudBase Inc (2021–Present)",
    "- Automated deployment pipelines for 15 microservices using GitHub Actions and ArgoCD",
    "- Integrated GitHub Copilot into team IDE setup; drove 30% adoption within 6 months",
    "- Built observability stack (Prometheus, Grafana, Loki)",
    "- Used OpenAI API (GPT-4) to build an internal Slack bot that answers infra FAQs",
    "",
    "Platform Engineer — StartupXYZ (2018–2021)",
    "- Managed Kubernetes clusters on GKE; implemented autoscaling and cost dashboards",
    "- Wrote Python scripts for log aggregation and alerting",
    "",
    "SKILLS",
    "Linux, Kubernetes, Terraform, AWS, GCP, GitHub Actions, Python, Bash, Docker, "
    "Prometheus, Grafana, GitHub Copilot, OpenAI API (basic usage)",
    "",
    "EDUCATION",
    "B.S. Information Systems — Cal State Fullerton, 2016",
    "",
    "CERTIFICATIONS",
    "AWS Solutions Architect – Associate (2022)",
    "CKA – Certified Kubernetes Administrator (2021)",
]

make_docx(OUT / "bob_martinez_devops.docx", BOB)
print("Created bob_martinez_devops.docx")


# ── Resume 3: Carol Nguyen — HR  (expected score ~1) ────────────────────

CAROL = """\
Carol Nguyen
HR Business Partner | carol.nguyen@email.com | LinkedIn: /in/carolnguyen

PROFESSIONAL SUMMARY
Experienced HR professional with 10 years in talent acquisition, employee relations,
and organizational development. Passionate about building inclusive workplace cultures.

EXPERIENCE

Senior HR Business Partner — Meridian Corp (2019–Present)
- Manages full-cycle recruitment for 200+ positions annually across 5 departments
- Conducts onboarding, performance review cycles, and manager training programs
- Uses Workday HRIS and LinkedIn Recruiter for candidate sourcing
- Partnered with IT to roll out Microsoft Copilot to the company (change management lead)

HR Generalist — BlueWave Services (2015–2019)
- Administered benefits, handled employee relations cases, supported L&D programs
- Coordinated annual engagement surveys and drafted action plans

SKILLS
Workday, ADP, LinkedIn Recruiter, Microsoft Office Suite, Microsoft Copilot (user),
Conflict resolution, Employment law, Onboarding design

EDUCATION
B.A. Psychology — University of Oregon, 2014
SHRM-CP Certification, 2017
"""

make_txt(OUT / "carol_nguyen_hr.txt", CAROL)
print("Created carol_nguyen_hr.txt")


# ── Resume 4: Diana Patel — Data Scientist  (expected score ~7) ──────────

DIANA = """\
Diana Patel
Data Scientist | diana.patel@email.com | github.com/dianapatel-ds

SUMMARY
Data scientist with 4 years of experience applying ML and NLP to business problems.
Recent focus on LLM-powered applications and prompt engineering for production use cases.

EXPERIENCE

Data Scientist II — RetailMind Analytics (2022–Present)
- Built customer churn prediction model (XGBoost, AUC 0.92) deployed to production via FastAPI
- Developed LLM-powered product review summarizer using Claude API with structured prompt chains
- Implemented semantic search using sentence-transformers + Pinecone for internal knowledge base
- Designed A/B testing framework for ML model evaluation across 50K daily users

Data Analyst → Data Scientist — Quentis Health (2020–2022)
- Transitioned from SQL analytics to building supervised classifiers for patient risk stratification
- Completed fast.ai Practical Deep Learning course; applied knowledge to NLP text classification

SKILLS
Python, pandas, scikit-learn, XGBoost, HuggingFace sentence-transformers, FastAPI,
Claude API, OpenAI API, Pinecone, SQL, dbt, Tableau, prompt engineering, A/B testing

EDUCATION
B.S. Statistics — University of Michigan, 2020
Coursera Machine Learning Specialization (Andrew Ng), 2021
"""

make_txt(OUT / "diana_patel_data_scientist.txt", DIANA)
print("Created diana_patel_data_scientist.txt")


# ── Resume 5: Evan Brooks — Junior Dev  (expected score ~2) ─────────────

EVAN = [
    "Evan Brooks",
    "Junior Software Developer | evan.brooks@email.com",
    "",
    "OBJECTIVE",
    "Recent CS graduate seeking a junior developer role. Eager to learn and contribute "
    "to a collaborative team environment.",
    "",
    "EDUCATION",
    "B.S. Computer Science — Arizona State University, May 2024",
    "GPA: 3.2",
    "",
    "EXPERIENCE",
    "Software Development Intern — LocalBiz App (Summer 2023)",
    "- Developed REST API endpoints in Node.js/Express for a small business platform",
    "- Fixed bugs in the React frontend and wrote unit tests with Jest",
    "- Used GitHub Copilot to assist with boilerplate code generation (first exposure to AI tools)",
    "",
    "Campus IT Help Desk — ASU (2022–2024)",
    "- Provided technical support to students and faculty",
    "- Troubleshot hardware, software, and network issues",
    "",
    "PROJECTS",
    "Personal Portfolio Website (React, deployed on Vercel)",
    "Todo App (Node.js + SQLite + HTML/CSS)",
    "",
    "SKILLS",
    "JavaScript, Node.js, React, HTML, CSS, SQL, Git, GitHub Copilot (basic user)",
    "",
    "CERTIFICATIONS",
    "Currently studying for AWS Cloud Practitioner",
]

make_docx(OUT / "evan_brooks_junior_dev.docx", EVAN)
print("Created evan_brooks_junior_dev.docx")

print("\nAll sample resumes created in sample_resumes/")
